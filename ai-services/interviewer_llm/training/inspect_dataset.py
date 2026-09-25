"""
Step 2 - Inspect the downloaded dataset and print a quality report.

    python training/inspect_dataset.py                 # uses data/dataset_path.txt
    python training/inspect_dataset.py --path <dir>    # or point at a folder / file

Read-only: it never modifies the original files. Nothing is assumed about column
names - the structure (which column is the question / answer / category /
difficulty) is DETECTED from the file, and the evidence is printed so you can
check it. Standard library only (optional: openpyxl for .xlsx, pyarrow for .parquet).

Outputs (besides stdout):
    data/reports/dataset_inspection.txt
    data/reports/dataset_inspection.json
"""
from __future__ import annotations

import argparse
import csv
import json
import random
import re
import sqlite3
import statistics
import sys
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
REPORT_DIR = ROOT / "data" / "reports"
PATH_FILE = ROOT / "data" / "dataset_path.txt"

MISSING_TOKENS = {"", "nan", "none", "null", "n/a", "na", "-", "--", "?", "tbd", "todo"}
ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
MOJIBAKE = re.compile(r"(â€|Ã.|Â |ï¿½|�)")
QUESTION_STARTS = (
    "what", "why", "how", "when", "where", "which", "who", "explain", "describe", "define", "compare", "differentiate",
    "list", "write", "discuss", "give", "state", "name", "can", "does", "do", "is", "are", "should", "implement",
    "design", "outline", "identify", "distinguish", "illustrate", "elaborate", "summarize", "mention", "tell",
)

# The subject areas the interviewer model is meant to cover, with keywords used to see what the data supports.
TARGET_TOPICS = {
    "Object-Oriented Programming": ["polymorphism", "inheritance", "encapsulation", "abstraction", "abstract class", "interface", "oop", "constructor", "overload", "overrid"],
    "DBMS / SQL": ["sql", "database", "dbms", "normalization", "join", "index", "transaction", "acid", "primary key", "foreign key", "query", "nosql"],
    "Operating Systems": ["operating system", "process", "thread", "deadlock", "scheduling", "paging", "virtual memory", "semaphore", "mutex", "kernel"],
    "Computer Networks": ["network", "tcp", "udp", "http", "dns", "ip address", "osi", "router", "socket", "protocol"],
    "SDLC / Agile": ["sdlc", "agile", "scrum", "waterfall", "sprint", "kanban", "requirements", "lifecycle"],
    "Testing": ["testing", "unit test", "integration test", "regression", "qa ", "test case", "tdd", "bug"],
    "Design / Architecture": ["design pattern", "singleton", "factory", "observer", "mvc", "microservice", "architecture", "solid", "scalab", "uml"],
    "Data Structures / Algorithms": ["array", "linked list", "stack", "queue", "tree", "graph", "hash", "sorting", "complexity", "big o", "recursion"],
    "Version control / DevOps": ["git", "version control", "ci/cd", "docker", "deployment", "devops", "branch"],
    "Programming languages": ["python", "java", "c++", "javascript", "compil", "interpret", "garbage collect", "memory", "pointer"],
    "Web / APIs / Security": ["api", "rest", "web", "authentication", "encryption", "security", "cors", "sql injection", "xss", "http"],
}

LINES: list[str] = []


def log(text: str = "") -> None:
    print(text)
    LINES.append(text)


def section(title: str) -> None:
    log("")
    log("=" * 78)
    log(title)
    log("=" * 78)


# ───────────────────────────── reading ─────────────────────────────

def read_text(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    for enc in ENCODINGS:
        try:
            return raw.decode(enc), enc
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8 (with replacement characters)"


def read_table(path: Path) -> tuple[list[dict], dict]:
    """Returns (rows, info). Rows are dicts of raw values; info describes how it was read."""
    ext = path.suffix.lower()
    info: dict = {"format": ext.lstrip(".") or "(none)"}

    if ext in {".csv", ".tsv", ".txt"}:
        text, enc = read_text(path)
        info["encoding"] = enc
        try:
            dialect = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel_tab if ext == ".tsv" else csv.excel
        info["delimiter"] = dialect.delimiter
        reader = csv.DictReader(text.splitlines(keepends=True) if False else _lines(text), dialect=dialect)
        rows = [dict(r) for r in reader]
        info["header"] = reader.fieldnames
        broken = sum(1 for r in rows if None in r)  # more fields than header
        short = sum(1 for r in rows if any(v is None for k, v in r.items() if k is not None))
        info["rows_with_extra_fields"] = broken
        info["rows_with_missing_fields"] = short
        for r in rows:
            r.pop(None, None)
        return rows, info

    if ext == ".jsonl":
        text, enc = read_text(path)
        info["encoding"] = enc
        rows, bad = [], 0
        for line in text.splitlines():
            if not line.strip():
                continue
            try:
                obj = json.loads(line)
                rows.append(obj if isinstance(obj, dict) else {"value": obj})
            except json.JSONDecodeError:
                bad += 1
        info["unparseable_lines"] = bad
        return rows, info

    if ext == ".json":
        text, enc = read_text(path)
        info["encoding"] = enc
        data = json.loads(text)
        if isinstance(data, dict):
            # {"data": [...]} / {"questions": [...]} / column-oriented dict of lists
            lists = {k: v for k, v in data.items() if isinstance(v, list)}
            if len(lists) == 1 and all(isinstance(x, dict) for x in next(iter(lists.values()))):
                key = next(iter(lists))
                info["root_key"] = key
                return list(lists[key]), info
            if lists and all(len(v) == len(next(iter(lists.values()))) for v in lists.values()):
                n = len(next(iter(lists.values())))
                info["orientation"] = "columns"
                return [{k: v[i] for k, v in lists.items()} for i in range(n)], info
            return [data], info
        return [x if isinstance(x, dict) else {"value": x} for x in data], info

    if ext in {".xlsx", ".xlsm"}:
        try:
            import openpyxl  # type: ignore
        except ImportError:
            info["skipped"] = "needs openpyxl (pip install openpyxl)"
            return [], info
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h) if h is not None else f"col{i}" for i, h in enumerate(next(it))]
        info["sheets"] = wb.sheetnames
        return [dict(zip(header, r)) for r in it], info

    if ext == ".parquet":
        try:
            import pyarrow.parquet as pq  # type: ignore
        except ImportError:
            info["skipped"] = "needs pyarrow (pip install pyarrow)"
            return [], info
        return pq.read_table(path).to_pylist(), info

    if ext in {".db", ".sqlite", ".sqlite3"}:
        con = sqlite3.connect(f"file:{path}?mode=ro", uri=True)
        tables = [r[0] for r in con.execute("select name from sqlite_master where type='table'")]
        info["tables"] = tables
        if not tables:
            return [], info
        cur = con.execute(f'select * from "{tables[0]}"')
        cols = [c[0] for c in cur.description]
        return [dict(zip(cols, r)) for r in cur.fetchall()], info

    info["skipped"] = "unsupported format (not a table)"
    return [], info


def _lines(text: str):
    # csv needs newline-preserving iteration so quoted multi-line fields survive
    import io

    return io.StringIO(text, newline="")


# ───────────────────────────── analysis helpers ─────────────────────────────

def is_missing(v) -> bool:
    return v is None or (isinstance(v, float) and v != v) or (isinstance(v, str) and v.strip().lower() in MISSING_TOKENS)


def norm(text: str) -> str:
    return " ".join(re.sub(r"[^a-z0-9\s]", " ", str(text).lower()).split())


def tokens(text: str) -> set[str]:
    return set(norm(text).split())


def guess_type(values: list) -> str:
    vals = [v for v in values if not is_missing(v)]
    if not vals:
        return "empty"
    def ok(fn):
        try:
            return all(fn(str(v).strip()) is not None for v in vals)
        except Exception:
            return False
    if ok(int):
        return "integer"
    if ok(float):
        return "float"
    if all(str(v).strip().lower() in {"true", "false", "yes", "no"} for v in vals):
        return "boolean"
    return "text"


def stats(nums: list[float]) -> str:
    if not nums:
        return "-"
    return f"min {min(nums):.0f} · median {statistics.median(nums):.0f} · mean {statistics.mean(nums):.0f} · max {max(nums):.0f}"


def detect_roles(cols: list[str], rows: list[dict]) -> dict:
    """Guess which column plays which part; returns roles plus the evidence used."""
    evidence: dict[str, str] = {}
    types = {c: guess_type([r.get(c) for r in rows]) for c in cols}
    n = max(len(rows), 1)
    uniq = {c: len({str(r.get(c)).strip() for r in rows if not is_missing(r.get(c))}) for c in cols}

    def by_name(pattern: str, exclude: str = "") -> list[str]:
        return [c for c in cols if re.search(pattern, c, re.I) and not (exclude and re.search(exclude, c, re.I))]

    roles: dict[str, str | None] = {"id": None, "question": None, "answer": None, "category": None, "difficulty": None}

    ids = by_name(r"number|(^|[^a-z])id($|[^a-z])|index|^no\.?$|serial|#")
    if ids:
        roles["id"] = ids[0]; evidence["id"] = f"name matches '{ids[0]}'"

    q = by_name(r"question|query|prompt|problem|title", exclude=r"number|id|index")
    if q:
        roles["question"] = q[0]; evidence["question"] = f"name matches '{q[0]}'"
    else:
        best = max((c for c in cols if types[c] == "text"), key=lambda c: sum(str(r.get(c) or "").strip().endswith("?") for r in rows), default=None)
        if best:
            share = sum(str(r.get(best) or "").strip().endswith("?") for r in rows) / n
            if share > 0.3:
                roles["question"] = best; evidence["question"] = f"{share:.0%} of its values end with '?'"

    a = by_name(r"answer|response|solution|explanation|reply|output")
    if a:
        roles["answer"] = a[0]; evidence["answer"] = f"name matches '{a[0]}'"
    else:
        txt = [c for c in cols if types[c] == "text" and c != roles["question"] and uniq[c] > 0.8 * n]
        if txt:
            best = max(txt, key=lambda c: statistics.mean(len(str(r.get(c) or "")) for r in rows))
            roles["answer"] = best; evidence["answer"] = "longest mostly-unique text column"

    d = by_name(r"difficult|level|complexity|hardness|grade")
    if d:
        roles["difficulty"] = d[0]; evidence["difficulty"] = f"name matches '{d[0]}'"

    cat = by_name(r"categor|topic|subject|domain|skill|area|section|tag|class|type|group", exclude=r"difficult|level")
    if cat:
        roles["category"] = cat[0]; evidence["category"] = f"name matches '{cat[0]}'"
    else:
        low = [c for c in cols if types[c] == "text" and 1 < uniq[c] <= max(30, n // 10) and c not in roles.values()]
        if low:
            roles["category"] = low[0]; evidence["category"] = f"low-cardinality text column ({uniq[low[0]]} values)"
    return {"roles": roles, "evidence": evidence, "types": types}


def near_duplicates(items: list[tuple[int, str]], threshold: float = 0.85, limit: int = 4000) -> list[tuple[int, int, float]]:
    """Pairs of (row, row, similarity) whose normalised text is very close (but not identical)."""
    data = [(i, norm(t), tokens(t)) for i, t in items if norm(t)]
    if len(data) > limit:
        random.Random(0).shuffle(data)
        data = data[:limit]
    pairs = []
    for a in range(len(data)):
        ia, na, ta = data[a]
        for b in range(a + 1, len(data)):
            ib, nb, tb = data[b]
            if na == nb:
                continue
            if abs(len(na) - len(nb)) > 0.4 * max(len(na), len(nb)):
                continue
            jac = len(ta & tb) / max(len(ta | tb), 1)
            if jac < 0.6:
                continue
            ratio = SequenceMatcher(None, na, nb).ratio()
            if ratio >= threshold:
                pairs.append((ia, ib, ratio))
    return sorted(pairs, key=lambda x: -x[2])


# ───────────────────────────── per-file report ─────────────────────────────

def inspect_file(path: Path, root: Path) -> dict:
    rel = path.relative_to(root) if path.is_relative_to(root) else path
    section(f"FILE: {rel}")
    size = path.stat().st_size
    log(f"size            : {size:,} bytes ({size / 1024:.1f} KB)")
    rows, info = read_table(path)
    for k, v in info.items():
        log(f"{k:<16}: {v}")
    result: dict = {"file": str(rel), "bytes": size, "info": {k: str(v) for k, v in info.items()}, "records": len(rows)}
    if info.get("skipped"):
        log(f"SKIPPED         : {info['skipped']}")
        return result
    if not rows:
        log("records         : 0")
        return result

    cols = list(dict.fromkeys(k for r in rows for k in r.keys()))
    log(f"records         : {len(rows)}")
    log(f"columns ({len(cols)})     : {cols}")
    result["columns"] = cols

    det = detect_roles(cols, rows)
    roles, types = det["roles"], det["types"]
    log("")
    log("Detected structure (from the data, not assumed):")
    for role, col in roles.items():
        log(f"  {role:<11}: {col!r:<22} {('- ' + det['evidence'][role]) if col else '- not found'}")
    result["roles"] = roles

    # ── columns
    log("")
    log("Columns:")
    log(f"  {'column':<22}{'type':<9}{'missing':>8}{'unique':>8}   text length (chars)")
    col_report = {}
    for c in cols:
        vals = [r.get(c) for r in rows]
        miss = sum(is_missing(v) for v in vals)
        present = [str(v) for v in vals if not is_missing(v)]
        uniq = len(set(present))
        lens = [len(v) for v in present] if types[c] == "text" else []
        log(f"  {c:<22}{types[c]:<9}{miss:>8}{uniq:>8}   {stats(lens)}")
        col_report[c] = {"type": types[c], "missing": miss, "unique": uniq}
    result["column_report"] = col_report

    # ── examples
    log("")
    log("Example records (first 3, then 3 random):")
    rng = random.Random(7)
    picks = list(range(min(3, len(rows)))) + rng.sample(range(min(3, len(rows)), len(rows)), min(3, max(len(rows) - 3, 0)))
    for i in picks:
        log(f"  [row {i + 1}]")
        for c in cols:
            v = str(rows[i].get(c)).replace("\n", " ")
            log(f"     {c}: {v[:160]}{'…' if len(v) > 160 else ''}")

    # ── whole-row duplicates
    keyed = [json.dumps(r, sort_keys=True, default=str) for r in rows]
    dup_rows = len(keyed) - len(set(keyed))
    log("")
    log(f"Exact duplicate rows                    : {dup_rows}")
    result["duplicate_rows"] = dup_rows

    qcol, acol, ccol, dcol = roles["question"], roles["answer"], roles["category"], roles["difficulty"]
    if not qcol:
        log("No question column detected - question-level checks skipped.")
        return result

    Q = [None if is_missing(r.get(qcol)) else str(r.get(qcol)).strip() for r in rows]
    A = [None if (not acol or is_missing(r.get(acol))) else str(r.get(acol)).strip() for r in rows]

    # ── question checks
    section(f"QUALITY: questions (column '{qcol}')")
    empty_q = [i for i, q in enumerate(Q) if not q]
    short_q = [i for i, q in enumerate(Q) if q and (len(q) < 15 or len(q.split()) < 4)]
    long_q = [i for i, q in enumerate(Q) if q and len(q) > 400]
    def question_like(q: str) -> bool:
        return q.endswith("?") or q.lower().split()[0].strip(",.:;") in QUESTION_STARTS
    not_q = [i for i, q in enumerate(Q) if q and not question_like(q)]
    mojibake = [i for i, q in enumerate(Q) if q and MOJIBAKE.search(q)]
    html = [i for i, q in enumerate(Q) if q and re.search(r"<[a-z/][^>]*>|https?://", q, re.I)]
    ws = [i for i, q in enumerate(Q) if q and rows[i].get(qcol) != q]
    qlens = [len(q) for q in Q if q]
    log(f"Empty / missing questions               : {len(empty_q)}")
    log(f"Very short (<15 chars or <4 words)      : {len(short_q)}")
    log(f"Very long (>400 chars)                  : {len(long_q)}")
    log(f"Not question-like (no '?' / no leading verb): {len(not_q)}")
    log(f"Ends with '?'                           : {sum(1 for q in Q if q and q.endswith('?'))} of {len(rows) - len(empty_q)}")
    log(f"Encoding damage (mojibake)              : {len(mojibake)}")
    log(f"HTML / URLs inside the text             : {len(html)}")
    log(f"Leading/trailing whitespace             : {len(ws)}")
    log(f"Question length (chars)                 : {stats(qlens)}")
    for label, idxs in (("not question-like", not_q), ("very short", short_q)):
        for i in idxs[:3]:
            log(f"     e.g. {label}: row {i + 1}: {Q[i]!r}")

    exact = defaultdict(list)
    normed = defaultdict(list)
    for i, q in enumerate(Q):
        if q:
            exact[q].append(i)
            normed[norm(q)].append(i)
    dup_exact = {k: v for k, v in exact.items() if len(v) > 1}
    dup_norm = {k: v for k, v in normed.items() if len(v) > 1}
    extra_exact = sum(len(v) - 1 for v in dup_exact.values())
    extra_norm = sum(len(v) - 1 for v in dup_norm.values())
    log("")
    log(f"Duplicate questions (exact text)        : {extra_exact} extra rows in {len(dup_exact)} groups")
    log(f"Duplicate questions (ignoring case/punct): {extra_norm} extra rows in {len(dup_norm)} groups")
    for v in list(dup_norm.values())[:3]:
        log(f"     e.g. rows {[i + 1 for i in v]}: {Q[v[0]]!r}")

    pairs = near_duplicates([(i, q) for i, q in enumerate(Q) if q])
    log(f"Near-duplicate pairs (similarity >= 0.85): {len(pairs)}")
    for a, b, s in pairs[:4]:
        log(f"     {s:.2f}  row {a + 1}: {Q[a]!r}")
        log(f"           row {b + 1}: {Q[b]!r}")
    result["question_quality"] = {
        "empty": len(empty_q), "very_short": len(short_q), "very_long": len(long_q), "not_question_like": len(not_q),
        "mojibake": len(mojibake), "html_or_url": len(html), "duplicate_exact_extra": extra_exact,
        "duplicate_normalised_extra": extra_norm, "near_duplicate_pairs": len(pairs),
    }

    # ── answer checks
    bad_a: set[int] = set()
    if acol:
        section(f"QUALITY: answers (column '{acol}')")
        empty_a = [i for i, a in enumerate(A) if not a]
        short_a = [i for i, a in enumerate(A) if a and len(a) < 20]
        long_a = [i for i, a in enumerate(A) if a and len(a) > 1500]
        same = [i for i, (q, a) in enumerate(zip(Q, A)) if q and a and norm(q) == norm(a)]
        junk = [i for i, a in enumerate(A) if a and re.search(r"\b(todo|tbd|lorem ipsum|i don'?t know|placeholder)\b", a, re.I)]
        moj_a = [i for i, a in enumerate(A) if a and MOJIBAKE.search(a)]
        alens = [len(a) for a in A if a]
        log(f"Empty / missing answers                 : {len(empty_a)}")
        log(f"Very short (<20 chars)                  : {len(short_a)}")
        log(f"Very long (>1500 chars)                 : {len(long_a)}")
        log(f"Answer identical to the question        : {len(same)}")
        log(f"Placeholder / junk answers              : {len(junk)}")
        log(f"Encoding damage (mojibake)              : {len(moj_a)}")
        log(f"Answer length (chars)                   : {stats(alens)}")
        log(f"Answer length (words)                   : {stats([len(a.split()) for a in A if a])}")
        bad_a = set(empty_a) | set(short_a) | set(same) | set(junk)
        # same question, different answers
        by_q = defaultdict(set)
        for q, a in zip(Q, A):
            if q and a:
                by_q[norm(q)].add(norm(a))
        conflicts = sum(1 for v in by_q.values() if len(v) > 1)
        by_a = defaultdict(set)
        for q, a in zip(Q, A):
            if q and a:
                by_a[norm(a)].add(norm(q))
        shared = sum(1 for v in by_a.values() if len(v) > 1)
        log(f"Same question, DIFFERENT answers        : {conflicts} groups")
        log(f"Same answer used for DIFFERENT questions: {shared} groups")
        result["answer_quality"] = {"empty": len(empty_a), "very_short": len(short_a), "very_long": len(long_a), "identical_to_question": len(same), "junk": len(junk), "conflicting_answers": conflicts}

    # ── categories / difficulty
    for label, col in (("category", ccol), ("difficulty", dcol)):
        if not col:
            continue
        section(f"DISTRIBUTION: {label} (column '{col}')")
        vals = Counter(str(r.get(col)).strip() if not is_missing(r.get(col)) else "(missing)" for r in rows)
        log(f"{len(vals)} distinct value(s)")
        for v, c in vals.most_common(40):
            log(f"  {c:>5}  {c / len(rows):>6.1%}  {v}")
        # near-identical labels (typos / case) are a common cleaning issue
        keys = defaultdict(list)
        for v in vals:
            keys[re.sub(r"[^a-z0-9]", "", v.lower())].append(v)
        variants = [v for v in keys.values() if len(v) > 1]
        if variants:
            log(f"  NOTE spelling/case variants of the same label: {variants}")
        result[f"{label}_values"] = dict(vals)
    if ccol and dcol:
        section("CROSS-TAB: category x difficulty")
        cats = sorted({str(r.get(ccol)).strip() for r in rows})
        diffs = sorted({str(r.get(dcol)).strip() for r in rows})
        log(f"  {'':<28}" + "".join(f"{d:>9}" for d in diffs))
        for cv in cats:
            log(f"  {cv[:27]:<28}" + "".join(f"{sum(1 for r in rows if str(r.get(ccol)).strip() == cv and str(r.get(dcol)).strip() == d):>9}" for d in diffs))

    # ── topic coverage against the subjects the model should handle
    section("TOPIC COVERAGE (keyword search over question + answer text)")
    blob = [(" ".join(filter(None, [Q[i], A[i]]))).lower() for i in range(len(rows))]
    log(f"  {'subject area':<32}{'matching records':>17}")
    cover = {}
    for topic, kws in TARGET_TOPICS.items():
        n_match = sum(1 for b in blob if any(k in b for k in kws))
        cover[topic] = n_match
        flag = "   <-- little or no data" if n_match < 5 else ""
        log(f"  {topic:<32}{n_match:>17}{flag}")
    result["topic_coverage"] = cover

    # ── usable estimate
    section("USABLE RECORD ESTIMATE")
    bad_q = set(empty_q) | set(short_q) | set(long_q) | set(not_q)
    seen, dup_drop = set(), set()
    for i, q in enumerate(Q):
        if not q:
            continue
        k = norm(q)
        if k in seen:
            dup_drop.add(i)
        seen.add(k)
    near_drop = {b for a, b, _ in pairs}
    ok_q = [i for i in range(len(rows)) if i not in bad_q and i not in dup_drop and i not in near_drop]
    ok_qa = [i for i in ok_q if (not acol) or i not in bad_a]
    log(f"Total records                                   : {len(rows)}")
    log(f"  removed: empty/short/long/not a question      : {len(bad_q)}")
    log(f"  removed: duplicate question (keep first)      : {len(dup_drop - bad_q)}")
    log(f"  removed: near-duplicate (keep first)          : {len(near_drop - bad_q - dup_drop)}")
    log(f"Usable for QUESTION generation                  : {len(ok_q)}")
    if acol:
        log(f"  removed additionally: bad/empty answer        : {len(set(ok_q) - set(ok_qa))}")
        log(f"Usable as QUESTION + ANSWER pairs               : {len(ok_qa)}")
    log(f"After an 80/10/10 split (question generation)   : train {int(len(ok_q) * .8)} / val {int(len(ok_q) * .1)} / test {len(ok_q) - int(len(ok_q) * .8) - int(len(ok_q) * .1)}")
    result["usable"] = {"total": len(rows), "usable_questions": len(ok_q), "usable_qa_pairs": len(ok_qa)}
    return result


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--path", help="dataset folder or file (default: location saved by download_dataset.py)")
    args = ap.parse_args()

    target = Path(args.path) if args.path else (Path(PATH_FILE.read_text(encoding="utf-8").strip()) if PATH_FILE.exists() else None)
    if not target or not target.exists():
        print("No dataset found. Run  python training/download_dataset.py  first (or pass --path).", file=sys.stderr)
        return 1

    files = [target] if target.is_file() else sorted(p for p in target.rglob("*") if p.is_file())
    root = target.parent if target.is_file() else target
    section("DATASET OVERVIEW")
    log(f"location : {target}")
    log(f"files    : {len(files)}   total {sum(f.stat().st_size for f in files):,} bytes")
    for f in files:
        log(f"  - {f.relative_to(root)}  ({f.stat().st_size:,} bytes)")

    results = [inspect_file(f, root) for f in files]

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    (REPORT_DIR / "dataset_inspection.txt").write_text("\n".join(LINES), encoding="utf-8")
    (REPORT_DIR / "dataset_inspection.json").write_text(json.dumps({"location": str(target), "files": results}, indent=2, default=str), encoding="utf-8")
    print(f"\nReport saved to {REPORT_DIR.relative_to(ROOT.parent)}/dataset_inspection.(txt|json)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
